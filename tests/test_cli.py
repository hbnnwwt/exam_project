import json
import os
import subprocess
import sys
from pathlib import Path
import tomllib

import openpyxl


def write_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="学号")
    ws.cell(row=2, column=1, value="参考答案")
    ws.cell(row=1, column=2, value=1)
    ws.cell(row=2, column=2, value="A")
    wb.save(path)
    wb.close()


def write_legacy_root(root: Path) -> None:
    (root / "config").mkdir(parents=True)
    (root / "config" / "sheet_layout.json").write_text(
        json.dumps(
            {
                "layout": {},
                "scoring": {},
                "choice": {
                    "question_start": 1,
                    "question_count": 1,
                    "options": ["A", "B"],
                },
            }
        ),
        encoding="utf-8",
    )
    write_xlsx(root / "参考答案.xlsx")


def run_cli(
    *args: str,
    env_overrides: dict[str, str] | None = None,
) -> subprocess.CompletedProcess[str]:
    env = os.environ.copy()
    repo_root = Path(__file__).resolve().parents[1]
    pythonpath = str(repo_root / "src")
    existing = env.get("PYTHONPATH")
    env["PYTHONPATH"] = pythonpath if not existing else os.pathsep.join([pythonpath, existing])
    if env_overrides:
        env.update(env_overrides)
    return subprocess.run(
        [sys.executable, "-m", "exam_project.cli", *args],
        text=True,
        capture_output=True,
        check=False,
        env=env,
    )


def test_cli_import_legacy_and_inspect(tmp_path: Path) -> None:
    legacy = tmp_path / "legacy"
    write_legacy_root(legacy)
    package_path = tmp_path / "demo.examproj"

    import_result = run_cli(
        "import-legacy",
        str(legacy),
        str(package_path),
        "--name",
        "演示项目",
    )
    assert import_result.returncode == 0
    assert package_path.exists()
    assert str(package_path) in import_result.stdout

    temp_root = tmp_path / "cli-temp"
    temp_root.mkdir()
    inspect_result = run_cli(
        "inspect",
        str(package_path),
        env_overrides={
            "TMP": str(temp_root),
            "TEMP": str(temp_root),
            "TMPDIR": str(temp_root),
        },
    )

    assert inspect_result.returncode == 0
    manifest = json.loads(inspect_result.stdout)
    assert manifest["name"] == "演示项目"
    assert manifest["assets"]["layout"] == "config/sheet_layout.json"
    assert not list(temp_root.glob(".exam_project_inspect_*.old-*"))


def test_cli_inspect_reports_project_errors_without_traceback(tmp_path: Path) -> None:
    package_path = tmp_path / "missing.examproj"

    result = run_cli("inspect", str(package_path))

    assert result.returncode == 2
    assert "Project package is not a valid ZIP file" in result.stderr
    assert "Traceback" not in result.stderr


def test_cli_import_legacy_reports_missing_assets_without_traceback(
    tmp_path: Path,
) -> None:
    legacy = tmp_path / "legacy"
    (legacy / "config").mkdir(parents=True)

    result = run_cli(
        "import-legacy",
        str(legacy),
        str(tmp_path / "demo.examproj"),
        "--name",
        "坏项目",
    )

    assert result.returncode == 2
    assert "缺少旧版布局文件" in result.stderr
    assert "Traceback" not in result.stderr


def test_console_script_entry_point_is_declared() -> None:
    pyproject = Path(__file__).resolve().parents[1] / "pyproject.toml"
    data = tomllib.loads(pyproject.read_text(encoding="utf-8"))

    assert data["project"]["scripts"]["exam-project"] == "exam_project.cli:main"
