import json
import zipfile
from pathlib import Path

import openpyxl
import pytest

from exam_project.core.errors import ProjectError, ProjectValidationError
from exam_project.core.legacy_import import import_legacy_project
from exam_project.core.package import ExamProjectPackage


def write_xlsx(path: Path, answers: dict[object, object] | None = None) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="学号")
    ws.cell(row=2, column=1, value="参考答案")
    data = answers if answers is not None else {1: "A"}
    for idx, (q_num, answer) in enumerate(data.items(), start=2):
        ws.cell(row=1, column=idx, value=q_num)
        ws.cell(row=2, column=idx, value=answer)
    wb.save(path)
    wb.close()


def write_legacy_root(
    root: Path,
    *,
    answers: dict[object, object] | None = None,
    include_baseline: bool = True,
) -> None:
    (root / "config").mkdir(parents=True)
    (root / "config" / "sheet_layout.json").write_text(
        json.dumps(
            {
                "layout": {},
                "scoring": {"choice_score": 1},
                "choice": {
                    "question_start": 1,
                    "question_count": 1,
                    "options": ["A", "B"],
                },
            }
        ),
        encoding="utf-8",
    )
    if include_baseline:
        (root / "config" / "blank_baseline.json").write_text("{}", encoding="utf-8")
    write_xlsx(root / "参考答案.xlsx", answers)


def test_import_legacy_project_creates_examproj(tmp_path: Path) -> None:
    legacy = tmp_path / "legacy"
    write_legacy_root(legacy)
    (legacy / "views").mkdir()
    (legacy / "views" / "designer_view.py").write_text("old ui", encoding="utf-8")
    (legacy / "modules").mkdir()
    (legacy / "modules" / "grading.py").write_text("old grading", encoding="utf-8")
    (legacy / "data" / "output").mkdir(parents=True)
    (legacy / "data" / "output" / "result.png").write_bytes(b"old output")
    package_path = tmp_path / "imported.examproj"

    import_legacy_project(legacy, package_path, name="导入项目")
    project = ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert project.manifest.name == "导入项目"
    assert project.layout_path.exists()
    assert project.answers_path.exists()
    assert project.baseline_path.exists()
    assert set(project.manifest.checksums) == {
        "design/answer_sheet.json",
        "config/sheet_layout.json",
        "answers/reference_answers.xlsx",
        "config/blank_baseline.json",
    }
    with zipfile.ZipFile(package_path, "r") as zf:
        names = set(zf.namelist())
    assert "views/designer_view.py" not in names
    assert "modules/grading.py" not in names
    assert "data/output/result.png" not in names


def test_import_legacy_project_allows_missing_baseline(tmp_path: Path) -> None:
    legacy = tmp_path / "legacy"
    write_legacy_root(legacy, include_baseline=False)
    package_path = tmp_path / "imported.examproj"

    import_legacy_project(legacy, package_path, name="无空白校准")
    project = ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert project.manifest.assets["baseline"] == "config/blank_baseline.json"
    assert not project.baseline_path.exists()
    assert set(project.manifest.checksums) == {
        "design/answer_sheet.json",
        "config/sheet_layout.json",
        "answers/reference_answers.xlsx",
    }


def test_import_legacy_project_rejects_missing_layout(tmp_path: Path) -> None:
    legacy = tmp_path / "legacy"
    write_legacy_root(legacy)
    (legacy / "config" / "sheet_layout.json").unlink()

    with pytest.raises(ProjectError, match="缺少旧版布局文件"):
        import_legacy_project(legacy, tmp_path / "imported.examproj", name="坏项目")


def test_import_legacy_project_rejects_missing_answers(tmp_path: Path) -> None:
    legacy = tmp_path / "legacy"
    write_legacy_root(legacy)
    (legacy / "参考答案.xlsx").unlink()

    with pytest.raises(ProjectError, match="缺少旧版参考答案"):
        import_legacy_project(legacy, tmp_path / "imported.examproj", name="坏项目")


def test_import_legacy_project_does_not_replace_existing_package_on_validation_error(
    tmp_path: Path,
) -> None:
    valid_legacy = tmp_path / "valid"
    write_legacy_root(valid_legacy)
    package_path = tmp_path / "imported.examproj"
    import_legacy_project(valid_legacy, package_path, name="原项目")

    invalid_legacy = tmp_path / "invalid"
    write_legacy_root(invalid_legacy, answers={1: "C"})

    with pytest.raises(ProjectValidationError) as exc:
        import_legacy_project(invalid_legacy, package_path, name="坏项目")

    opened = ExamProjectPackage.open(package_path, tmp_path / "opened")
    assert exc.value.code == "answer_layout_mismatch"
    assert opened.manifest.name == "原项目"
