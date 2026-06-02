from pathlib import Path

import openpyxl
import pytest

from exam_project.core.errors import ProjectValidationError
from exam_project.core.validation import classify_question, validate_answer_workbook


def make_xlsx(path: Path, answers: dict[object, object]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="student_id")
    ws.cell(row=2, column=1, value="answer")
    for idx, (q_num, answer) in enumerate(answers.items(), start=2):
        ws.cell(row=1, column=idx, value=q_num)
        ws.cell(row=2, column=idx, value=answer)
    wb.save(path)
    wb.close()


def test_classify_question_uses_choice_and_judge_ranges() -> None:
    layout = {
        "choice": {"question_start": 1, "question_count": 3, "options": ["A", "B"]},
        "judge": {"question_start": 4, "question_count": 2, "options": ["T", "F"]},
    }

    assert classify_question(1, layout) == "choice"
    assert classify_question(3, layout) == "choice"
    assert classify_question(4, layout) == "judge"
    assert classify_question(5, layout) == "judge"
    assert classify_question(6, layout) == "essay"


def test_validate_answer_workbook_accepts_matching_layout_options(
    tmp_path: Path,
) -> None:
    layout = {
        "choice": {"question_start": 1, "question_count": 2, "options": ["A", "B"]},
        "judge": {"question_start": 3, "question_count": 1, "options": ["T", "F"]},
    }
    path = tmp_path / "answers.xlsx"
    make_xlsx(path, {1: "A", 2: "B", 3: "T", 9: "free form"})

    validate_answer_workbook(path, layout)


def test_validate_answer_workbook_rejects_answer_layout_mismatch(
    tmp_path: Path,
) -> None:
    layout = {
        "choice": {"question_start": 1, "question_count": 2, "options": ["A", "B"]},
        "judge": {"question_start": 3, "question_count": 1, "options": ["T", "F"]},
    }
    path = tmp_path / "answers.xlsx"
    make_xlsx(path, {1: "A", 2: "C", 3: "T"})

    with pytest.raises(ProjectValidationError) as exc:
        validate_answer_workbook(path, layout)

    assert exc.value.code == "answer_layout_mismatch"
    assert "question 2" in exc.value.user_message
    assert "answer C" in exc.value.user_message
    assert "choice" in exc.value.user_message
    assert "A, B" in exc.value.user_message


def test_validate_answer_workbook_rejects_invalid_question_header(
    tmp_path: Path,
) -> None:
    layout = {
        "choice": {"question_start": 1, "question_count": 1, "options": ["A"]},
    }
    path = tmp_path / "answers.xlsx"
    make_xlsx(path, {"Q1": "A"})

    with pytest.raises(ProjectValidationError) as exc:
        validate_answer_workbook(path, layout)

    assert exc.value.code == "invalid_question_number"
    assert "Q1" in exc.value.user_message
