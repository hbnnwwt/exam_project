import json
from pathlib import Path

import openpyxl
import pytest

from exam_project.core.errors import ProjectValidationError
from exam_project.core.package import ExamProjectPackage


def make_xlsx(path: Path, answers: dict[object, object]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="student_id")
    ws.cell(row=2, column=1, value="answer")
    for idx, (q_num, answer) in enumerate(answers.items(), start=2):
        ws.cell(row=1, column=idx, value=q_num)
        ws.cell(row=2, column=idx, value=answer)
    wb.save(path)
    wb.close()


def write_project(
    root: Path,
    *,
    layout: dict | None = None,
    answers: dict[object, object] | None = None,
    checksums: dict[str, str] | None = None,
    include_baseline: bool = False,
) -> None:
    (root / "design").mkdir(exist_ok=True)
    (root / "config").mkdir(exist_ok=True)
    (root / "answers").mkdir(exist_ok=True)
    (root / "design" / "answer_sheet.json").write_text("{}", encoding="utf-8")
    (root / "config" / "sheet_layout.json").write_text(
        json.dumps(
            layout
            if layout is not None
            else {
                "choice": {
                    "question_start": 1,
                    "question_count": 1,
                    "options": ["A", "B"],
                }
            }
        ),
        encoding="utf-8",
    )
    make_xlsx(
        root / "answers" / "reference_answers.xlsx",
        answers if answers is not None else {1: "A"},
    )

    assets = {
        "design": "design/answer_sheet.json",
        "layout": "config/sheet_layout.json",
        "answers": "answers/reference_answers.xlsx",
    }
    if include_baseline:
        assets["baseline"] = "config/blank_baseline.json"

    (root / "project.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "project_id": "project-1",
                "name": "Final Exam",
                "created_at": "2026-06-02T20:00:00+08:00",
                "updated_at": "2026-06-02T20:00:00+08:00",
                "assets": assets,
                "exam": {"student_id_digits": 10, "question_types": ["choice"]},
                "checksums": checksums or {},
            }
        ),
        encoding="utf-8",
    )


def test_open_validates_project_before_replacing_target(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_project(workdir, answers={1: "C"})
    package_path = tmp_path / "bad.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    target_dir = tmp_path / "opened"
    target_dir.mkdir()
    (target_dir / "marker.txt").write_text("keep", encoding="utf-8")

    with pytest.raises(ProjectValidationError) as exc:
        ExamProjectPackage.open(package_path, target_dir)

    assert exc.value.code == "answer_layout_mismatch"
    assert (target_dir / "marker.txt").read_text(encoding="utf-8") == "keep"
    assert not (target_dir / "project.json").exists()


def test_open_rejects_missing_required_asset(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_project(workdir)
    (workdir / "config" / "sheet_layout.json").unlink()
    package_path = tmp_path / "bad.examproj"
    ExamProjectPackage.pack(workdir, package_path)

    with pytest.raises(ProjectValidationError) as exc:
        ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert exc.value.code == "missing_asset"


def test_open_rejects_checksum_mismatch(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_project(
        workdir,
        checksums={"design/answer_sheet.json": "sha256:" + "0" * 64},
    )
    package_path = tmp_path / "bad.examproj"
    ExamProjectPackage.pack(workdir, package_path)

    with pytest.raises(ProjectValidationError) as exc:
        ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert exc.value.code == "checksum_mismatch"


def test_open_rejects_missing_checksum_asset(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_project(
        workdir,
        checksums={"config/blank_baseline.json": "sha256:" + "0" * 64},
    )
    package_path = tmp_path / "bad.examproj"
    ExamProjectPackage.pack(workdir, package_path)

    with pytest.raises(ProjectValidationError) as exc:
        ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert exc.value.code == "missing_checksum_asset"


def test_open_allows_declared_missing_baseline(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_project(workdir, include_baseline=True)
    package_path = tmp_path / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)

    opened = ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert opened.manifest.assets["baseline"] == "config/blank_baseline.json"
    assert not (opened.workdir / "config" / "blank_baseline.json").exists()


def test_open_rejects_invalid_layout_json(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_project(workdir)
    (workdir / "config" / "sheet_layout.json").write_text("{", encoding="utf-8")
    package_path = tmp_path / "bad.examproj"
    ExamProjectPackage.pack(workdir, package_path)

    with pytest.raises(ProjectValidationError) as exc:
        ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert exc.value.code == "invalid_layout_json"
