import json
import shutil
import struct
import zipfile
import zlib
from pathlib import Path

import openpyxl
import pytest

import exam_project.core.package as package_module
from exam_project.core.errors import ProjectPackageError
from exam_project.core.package import ExamProjectPackage, safe_extract_zip


def write_reference_answers(path: Path, answers: dict[object, object]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="student_id")
    ws.cell(row=2, column=1, value="answer")
    for idx, (q_num, answer) in enumerate(answers.items(), start=2):
        ws.cell(row=1, column=idx, value=q_num)
        ws.cell(row=2, column=idx, value=answer)
    wb.save(path)
    wb.close()


def write_minimal_workdir(root: Path, name: str = "Final Exam") -> None:
    (root / "design").mkdir(exist_ok=True)
    (root / "config").mkdir(exist_ok=True)
    (root / "answers").mkdir(exist_ok=True)
    (root / "design" / "answer_sheet.json").write_text("{}", encoding="utf-8")
    (root / "config" / "sheet_layout.json").write_text(
        json.dumps(
            {
                "choice": {
                    "question_start": 1,
                    "question_count": 1,
                    "options": ["A", "B"],
                }
            }
        ),
        encoding="utf-8",
    )
    write_reference_answers(root / "answers" / "reference_answers.xlsx", {1: "A"})
    (root / "project.json").write_text(
        json.dumps(
            {
                "schema_version": 1,
                "project_id": "project-1",
                "name": name,
                "created_at": "2026-06-02T20:00:00+08:00",
                "updated_at": "2026-06-02T20:00:00+08:00",
                "assets": {
                    "design": "design/answer_sheet.json",
                    "layout": "config/sheet_layout.json",
                    "answers": "answers/reference_answers.xlsx",
                    "baseline": "config/blank_baseline.json",
                },
                "exam": {"student_id_digits": 10, "question_types": ["choice"]},
                "checksums": {},
            }
        ),
        encoding="utf-8",
    )


def write_raw_zip_entry(package_path: Path, raw_name: bytes, data: bytes = b"bad") -> None:
    crc = zlib.crc32(data) & 0xFFFFFFFF
    local = (
        struct.pack(
            "<IHHHHHIIIHH",
            0x04034B50,
            20,
            0,
            0,
            0,
            0,
            crc,
            len(data),
            len(data),
            len(raw_name),
            0,
        )
        + raw_name
        + data
    )
    central_offset = len(local)
    central = struct.pack(
        "<IHHHHHHIIIHHHHHII",
        0x02014B50,
        20,
        20,
        0,
        0,
        0,
        0,
        crc,
        len(data),
        len(data),
        len(raw_name),
        0,
        0,
        0,
        0,
        0,
        central_offset,
    ) + raw_name
    end = struct.pack(
        "<IHHHHIIH",
        0x06054B50,
        0,
        0,
        1,
        1,
        len(central),
        central_offset,
        0,
    )
    package_path.write_bytes(local + central + end)


def test_pack_and_open_project(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = tmp_path / "sample.examproj"

    ExamProjectPackage.pack(workdir, package_path)
    opened = ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert opened.package_path == package_path
    assert opened.workdir == tmp_path / "opened"
    assert opened.manifest.name == "Final Exam"
    assert (opened.workdir / "config" / "sheet_layout.json").exists()


def test_safe_extract_zip_rejects_path_traversal(tmp_path: Path) -> None:
    package_path = tmp_path / "evil.examproj"
    with zipfile.ZipFile(package_path, "w") as zf:
        zf.writestr("../evil.txt", "bad")

    with pytest.raises(ProjectPackageError):
        safe_extract_zip(package_path, tmp_path / "out")

    assert not (tmp_path / "evil.txt").exists()


def test_safe_extract_zip_rejects_raw_backslash_entry(tmp_path: Path) -> None:
    package_path = tmp_path / "evil.examproj"
    write_raw_zip_entry(package_path, b"dir\\evil.txt")

    with pytest.raises(ProjectPackageError):
        safe_extract_zip(package_path, tmp_path / "out")

    assert not (tmp_path / "out" / "dir" / "evil.txt").exists()


def test_safe_extract_zip_rejects_duplicate_entries(tmp_path: Path) -> None:
    package_path = tmp_path / "evil.examproj"
    with zipfile.ZipFile(package_path, "w") as zf:
        zf.writestr("project.json", "{}")
        with pytest.warns(UserWarning, match="Duplicate name"):
            zf.writestr("project.json", "[]")

    with pytest.raises(ProjectPackageError):
        safe_extract_zip(package_path, tmp_path / "out")


def test_safe_extract_zip_rejects_casefold_duplicate_entries(tmp_path: Path) -> None:
    package_path = tmp_path / "evil.examproj"
    with zipfile.ZipFile(package_path, "w") as zf:
        zf.writestr("Project.json", "{}")
        zf.writestr("project.json", "[]")

    with pytest.raises(ProjectPackageError):
        safe_extract_zip(package_path, tmp_path / "out")


@pytest.mark.parametrize(
    "entry_name",
    [
        "/absolute.txt",
        "C:/absolute.txt",
        "config/",
        "config//sheet_layout.json",
        "./config/sheet_layout.json",
        "config/./sheet_layout.json",
    ],
)
def test_safe_extract_zip_rejects_unsafe_names(tmp_path: Path, entry_name: str) -> None:
    package_path = tmp_path / "evil.examproj"
    with zipfile.ZipFile(package_path, "w") as zf:
        zf.writestr(entry_name, "bad")

    with pytest.raises(ProjectPackageError):
        safe_extract_zip(package_path, tmp_path / "out")


def test_save_creates_backup_and_replaces_with_verified_package(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir, name="Original Exam")
    package_path = tmp_path / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)

    write_minimal_workdir(workdir, name="Updated Exam")
    ExamProjectPackage.save(workdir, package_path)

    opened = ExamProjectPackage.open(package_path, tmp_path / "opened")
    backup = package_path.with_suffix(".examproj.bak")

    assert opened.manifest.name == "Updated Exam"
    assert backup.exists()
    assert ExamProjectPackage.open(backup, tmp_path / "backup").manifest.name == "Original Exam"
    assert not package_path.with_suffix(".examproj.tmp").exists()


def test_save_excludes_package_artifacts_when_package_is_inside_workdir(
    tmp_path: Path,
) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir, name="Original Exam")
    package_path = workdir / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    package_path.with_suffix(".examproj.bak").write_bytes(b"stale backup")
    package_path.with_suffix(".examproj.bak.tmp").write_bytes(b"stale backup tmp")

    write_minimal_workdir(workdir, name="Updated Exam")
    ExamProjectPackage.save(workdir, package_path)

    with zipfile.ZipFile(package_path, "r") as zf:
        names = set(zf.namelist())

    assert "sample.examproj" not in names
    assert "sample.examproj.bak" not in names
    assert "sample.examproj.bak.tmp" not in names


def test_pack_excludes_package_artifacts_by_default(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = workdir / "sample.examproj"
    package_path.write_bytes(b"old package")
    package_path.with_suffix(".examproj.tmp").write_bytes(b"old tmp")
    package_path.with_suffix(".examproj.bak").write_bytes(b"old backup")
    package_path.with_suffix(".examproj.bak.tmp").write_bytes(b"old backup tmp")

    ExamProjectPackage.pack(workdir, package_path)

    with zipfile.ZipFile(package_path, "r") as zf:
        names = set(zf.namelist())

    assert "sample.examproj" not in names
    assert "sample.examproj.tmp" not in names
    assert "sample.examproj.bak" not in names
    assert "sample.examproj.bak.tmp" not in names


def test_pack_skips_file_symlinks(tmp_path: Path) -> None:
    outside = tmp_path / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    link = workdir / "outside-link.txt"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("file symlinks are not available on this platform")
    package_path = tmp_path / "sample.examproj"

    ExamProjectPackage.pack(workdir, package_path)

    with zipfile.ZipFile(package_path, "r") as zf:
        assert "outside-link.txt" not in set(zf.namelist())


def test_open_does_not_replace_existing_target_when_package_is_invalid(
    tmp_path: Path,
) -> None:
    package_path = tmp_path / "invalid.examproj"
    with zipfile.ZipFile(package_path, "w") as zf:
        zf.writestr("not_project.json", "{}")
    target_dir = tmp_path / "opened"
    target_dir.mkdir()
    (target_dir / "marker.txt").write_text("keep", encoding="utf-8")

    with pytest.raises(ProjectPackageError):
        ExamProjectPackage.open(package_path, target_dir)

    assert (target_dir / "marker.txt").read_text(encoding="utf-8") == "keep"
    assert not (target_dir / "not_project.json").exists()


def test_open_rejects_package_parent_as_target(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = tmp_path / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    marker = tmp_path / "marker.txt"
    marker.write_text("keep", encoding="utf-8")

    with pytest.raises(ProjectPackageError):
        ExamProjectPackage.open(package_path, tmp_path)

    assert marker.read_text(encoding="utf-8") == "keep"


def test_open_rejects_package_ancestor_as_target(tmp_path: Path) -> None:
    target_dir = tmp_path / "opened"
    package_dir = target_dir / "packages"
    package_dir.mkdir(parents=True)
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = package_dir / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    marker = target_dir / "marker.txt"
    marker.write_text("keep", encoding="utf-8")

    with pytest.raises(ProjectPackageError):
        ExamProjectPackage.open(package_path, target_dir)

    assert marker.read_text(encoding="utf-8") == "keep"
    assert package_path.exists()


def test_open_allows_package_sibling_target(tmp_path: Path) -> None:
    package_dir = tmp_path / "packages"
    package_dir.mkdir()
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = package_dir / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    target_dir = tmp_path / "opened"

    opened = ExamProjectPackage.open(package_path, target_dir)

    assert opened.workdir == target_dir
    assert opened.manifest.name == "Final Exam"
    assert package_path.exists()


def test_open_restores_existing_target_when_replace_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = tmp_path / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    target_dir = tmp_path / "opened"
    target_dir.mkdir()
    (target_dir / "marker.txt").write_text("keep", encoding="utf-8")
    real_move = shutil.move
    moves = 0

    def fail_second_move(src: str, dst: str):
        nonlocal moves
        moves += 1
        if moves == 2:
            Path(dst).mkdir(parents=True, exist_ok=True)
            (Path(dst) / "partial.txt").write_text("partial", encoding="utf-8")
            raise OSError("simulated move failure")
        return real_move(src, dst)

    monkeypatch.setattr(shutil, "move", fail_second_move)

    with pytest.raises(OSError):
        ExamProjectPackage.open(package_path, target_dir)

    assert (target_dir / "marker.txt").read_text(encoding="utf-8") == "keep"
    assert not (target_dir / "partial.txt").exists()


def test_open_preserves_old_target_when_restore_move_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = tmp_path / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    target_dir = tmp_path / "opened"
    target_dir.mkdir()
    (target_dir / "marker.txt").write_text("keep", encoding="utf-8")
    real_move = shutil.move
    moves = 0

    def fail_restore_move(src: str, dst: str):
        nonlocal moves
        moves += 1
        if moves == 2:
            Path(dst).mkdir(parents=True, exist_ok=True)
            (Path(dst) / "partial.txt").write_text("partial", encoding="utf-8")
            raise OSError("simulated move failure")
        if moves == 3:
            raise OSError("simulated restore failure")
        return real_move(src, dst)

    monkeypatch.setattr(shutil, "move", fail_restore_move)

    with pytest.raises(OSError):
        ExamProjectPackage.open(package_path, target_dir)

    old_dirs = list(tmp_path.glob(".opened.old-*"))
    assert len(old_dirs) == 1
    assert (old_dirs[0] / "opened" / "marker.txt").read_text(encoding="utf-8") == "keep"
    assert not target_dir.exists()


def test_open_preserves_old_target_when_partial_cleanup_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir)
    package_path = tmp_path / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)
    target_dir = tmp_path / "opened"
    target_dir.mkdir()
    (target_dir / "marker.txt").write_text("keep", encoding="utf-8")
    real_move = shutil.move
    real_remove_existing_path = package_module._remove_existing_path
    moves = 0

    def fail_second_move(src: str, dst: str):
        nonlocal moves
        moves += 1
        if moves == 2:
            Path(dst).mkdir(parents=True, exist_ok=True)
            (Path(dst) / "partial.txt").write_text("partial", encoding="utf-8")
            raise OSError("simulated move failure")
        return real_move(src, dst)

    def fail_target_cleanup(path: Path) -> None:
        if path == target_dir:
            raise OSError("simulated cleanup failure")
        return real_remove_existing_path(path)

    monkeypatch.setattr(shutil, "move", fail_second_move)
    monkeypatch.setattr(package_module, "_remove_existing_path", fail_target_cleanup)

    with pytest.raises(OSError):
        ExamProjectPackage.open(package_path, target_dir)

    old_dirs = list(tmp_path.glob(".opened.old-*"))
    assert len(old_dirs) == 1
    assert (old_dirs[0] / "opened" / "marker.txt").read_text(encoding="utf-8") == "keep"
    assert (target_dir / "partial.txt").read_text(encoding="utf-8") == "partial"


def test_save_does_not_replace_existing_package_when_verify_fails(tmp_path: Path) -> None:
    workdir = tmp_path / "work"
    workdir.mkdir()
    write_minimal_workdir(workdir, name="Original Exam")
    package_path = tmp_path / "sample.examproj"
    ExamProjectPackage.pack(workdir, package_path)

    (workdir / "project.json").unlink()

    with pytest.raises(ProjectPackageError):
        ExamProjectPackage.save(workdir, package_path)

    assert ExamProjectPackage.open(package_path, tmp_path / "opened").manifest.name == "Original Exam"
    assert not package_path.with_suffix(".examproj.tmp").exists()
