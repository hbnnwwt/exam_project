from pathlib import Path

import openpyxl
import pytest

from exam_project.core.errors import ProjectValidationError
from exam_project.core.new_project import create_exam_project
from exam_project.core.package import ExamProjectPackage


def test_create_exam_project_creates_openable_blank_package(tmp_path: Path) -> None:
    package_path = tmp_path / "new.examproj"

    create_exam_project(package_path, name="新考试", student_id_digits=12)
    project = ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert project.manifest.name == "新考试"
    assert project.manifest.exam["student_id_digits"] == 12
    assert project.layout_path.exists()
    assert project.answers_path.exists()
    assert not project.baseline_path.exists()
    assert set(project.manifest.checksums) == {
        "design/answer_sheet.json",
        "config/sheet_layout.json",
        "answers/reference_answers.xlsx",
    }
    layout = project.load_layout()
    assert layout["student_id"]["digit_count"] == 12

    wb = openpyxl.load_workbook(project.answers_path, read_only=True, data_only=True)
    try:
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "学号"
        assert ws.cell(row=2, column=1).value == "参考答案"
        assert ws.max_column == 1
    finally:
        wb.close()


def test_create_exam_project_does_not_replace_existing_package_when_invalid(
    tmp_path: Path,
) -> None:
    package_path = tmp_path / "new.examproj"
    create_exam_project(package_path, name="原项目")

    with pytest.raises(ProjectValidationError):
        create_exam_project(package_path, name="")

    project = ExamProjectPackage.open(package_path, tmp_path / "opened")
    assert project.manifest.name == "原项目"


@pytest.mark.parametrize("student_id_digits", [0, -1, True, "10"])
def test_create_exam_project_rejects_invalid_student_id_digits(
    tmp_path: Path,
    student_id_digits: object,
) -> None:
    with pytest.raises(ProjectValidationError) as exc:
        create_exam_project(
            tmp_path / "new.examproj",
            name="坏项目",
            student_id_digits=student_id_digits,  # type: ignore[arg-type]
        )

    assert exc.value.code == "invalid_student_id_digits"
