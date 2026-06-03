import json
from pathlib import Path

import openpyxl
import pytest

from exam_project.core.errors import ProjectPackageError
from exam_project.core.new_project import create_exam_project
from exam_project.core.package import ExamProjectPackage
from exam_project.gui.session import (
    create_and_open_project,
    open_project,
    save_project,
    save_project_as,
    workspace_dir_for,
)


def _add_answer(project_package: Path, opened_dir: Path) -> None:
    project = ExamProjectPackage.open(project_package, opened_dir)
    layout = project.load_layout()
    layout["choice"] = {
        "question_start": 1,
        "question_count": 1,
        "options": ["A", "B"],
    }
    project.layout_path.write_text(
        json.dumps(layout, ensure_ascii=False),
        encoding="utf-8",
    )
    wb = openpyxl.load_workbook(project.answers_path)
    try:
        ws = wb.active
        ws.cell(row=1, column=2, value=1)
        ws.cell(row=2, column=2, value="A")
        wb.save(project.answers_path)
    finally:
        wb.close()
    save_project(
        open_project(project_package, workspace_root=opened_dir.parent / "workspace")
    )


def test_create_and_open_project_sets_stable_workspace(tmp_path: Path) -> None:
    package_path = tmp_path / "new.examproj"
    workspace_root = tmp_path / "workspace"

    session = create_and_open_project(
        package_path,
        name="新考试",
        student_id_digits=12,
        workspace_root=workspace_root,
    )

    assert session.package_path == package_path.resolve()
    assert session.workspace_root == workspace_root.resolve()
    assert session.workdir == workspace_dir_for(package_path, workspace_root)
    assert session.manifest.name == "新考试"
    assert session.project.answers_path.exists()
    assert session.dirty is False


def test_open_project_replaces_workspace_with_package_contents(tmp_path: Path) -> None:
    package_path = tmp_path / "sample.examproj"
    create_exam_project(package_path, name="原项目")
    workspace_root = tmp_path / "workspace"
    target = workspace_dir_for(package_path, workspace_root)
    target.mkdir(parents=True)
    (target / "stale.txt").write_text("stale", encoding="utf-8")

    session = open_project(package_path, workspace_root=workspace_root)

    assert session.manifest.name == "原项目"
    assert not (target / "stale.txt").exists()
    assert (target / "project.json").is_file()


def test_save_project_refreshes_manifest_checksum_and_backup(tmp_path: Path) -> None:
    package_path = tmp_path / "sample.examproj"
    session = create_and_open_project(package_path, name="原项目")
    original_updated_at = session.manifest.updated_at
    layout = session.project.load_layout()
    layout["choice"] = {
        "question_start": 1,
        "question_count": 1,
        "options": ["A", "B"],
    }
    session.project.layout_path.write_text(
        json.dumps(layout, ensure_ascii=False),
        encoding="utf-8",
    )
    wb = openpyxl.load_workbook(session.project.answers_path)
    try:
        ws = wb.active
        ws.cell(row=1, column=2, value=1)
        ws.cell(row=2, column=2, value="A")
        wb.save(session.project.answers_path)
    finally:
        wb.close()

    saved = save_project(session)
    reopened = ExamProjectPackage.open(package_path, tmp_path / "opened")

    assert saved.dirty is False
    assert saved.manifest.updated_at >= original_updated_at
    assert reopened.load_layout()["choice"]["question_count"] == 1
    assert package_path.with_suffix(".examproj.bak").exists()


def test_save_project_as_switches_package_path(tmp_path: Path) -> None:
    original = tmp_path / "sample.examproj"
    target = tmp_path / "copy.examproj"
    session = create_and_open_project(original, name="原项目")

    saved = save_project_as(session, target)

    assert saved.package_path == target.resolve()
    assert target.exists()
    assert original.exists()
    assert ExamProjectPackage.open(target, tmp_path / "opened").manifest.name == "原项目"


def test_bad_open_does_not_replace_existing_session_workspace(tmp_path: Path) -> None:
    package_path = tmp_path / "sample.examproj"
    session = create_and_open_project(package_path, name="原项目")
    marker = session.workdir / "marker.txt"
    marker.write_text("keep", encoding="utf-8")
    bad_package = tmp_path / "bad.examproj"
    bad_package.write_bytes(b"not zip")

    with pytest.raises(ProjectPackageError):
        open_project(bad_package, workspace_root=session.workspace_root)

    assert marker.read_text(encoding="utf-8") == "keep"
