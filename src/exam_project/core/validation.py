from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, LargeZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from exam_project.core.errors import ProjectValidationError


QUESTION_TYPES = ("choice", "judge", "essay")


def _int_or_none(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


TYPE_LABELS = {
    "choice": "选择题",
    "judge": "判断题",
    "essay": "主观题",
}


def classify_question(q_num: int, layout: dict) -> str | None:
    for q_type in QUESTION_TYPES:
        cfg = layout.get(q_type)
        if not isinstance(cfg, dict):
            continue
        start = _int_or_none(cfg.get("question_start"))
        count = _int_or_none(cfg.get("question_count"))
        if start is None or count is None or count <= 0:
            continue
        if start <= q_num < start + count:
            return q_type
    return None


def _options_for(q_type: str, layout: dict) -> list[str] | None:
    cfg = layout.get(q_type)
    if not isinstance(cfg, dict):
        return None
    options = cfg.get("options")
    if options is None:
        return None
    if isinstance(options, str):
        normalized_text = options.strip()
        if not normalized_text:
            raise ProjectValidationError(
                f"答题卡布局的 {TYPE_LABELS.get(q_type, q_type)} 选项无效",
                code="invalid_layout_options",
            )
        return [normalized_text]
    if isinstance(options, Mapping) or not isinstance(options, Sequence):
        raise ProjectValidationError(
            f"答题卡布局的 {TYPE_LABELS.get(q_type, q_type)} 选项无效",
            code="invalid_layout_options",
        )
    normalized = [str(option).strip() for option in options if not _is_blank(option)]
    if not normalized:
        raise ProjectValidationError(
            f"答题卡布局的 {TYPE_LABELS.get(q_type, q_type)} 选项无效",
            code="invalid_layout_options",
        )
    return normalized


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _question_number(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError(value)
    if isinstance(value, int):
        q_num = value
    elif isinstance(value, float) and value.is_integer():
        q_num = int(value)
    elif isinstance(value, str) and value.strip().isdigit():
        q_num = int(value.strip())
    else:
        raise ValueError(value)
    if q_num <= 0:
        raise ValueError(value)
    return q_num


def validate_answer_workbook(path: Path, layout: dict) -> None:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (
        OSError,
        BadZipFile,
        LargeZipFile,
        InvalidFileException,
        KeyError,
        ValueError,
        TypeError,
    ) as exc:
        raise ProjectValidationError(
            f"无法打开参考答案工作簿: {path}",
            code="invalid_answer_workbook",
        ) from exc

    try:
        ws = wb.active
        for col in range(2, ws.max_column + 1):
            q_raw = ws.cell(row=1, column=col).value
            answer_raw = ws.cell(row=2, column=col).value
            if _is_blank(q_raw) and _is_blank(answer_raw):
                continue
            if _is_blank(q_raw):
                raise ProjectValidationError(
                    f"参考答案第 {col} 列缺少题号",
                    code="missing_question_number",
                )
            if _is_blank(answer_raw):
                raise ProjectValidationError(
                    f"第 {q_raw} 题缺少参考答案",
                    code="missing_answer",
                )

            try:
                q_num = _question_number(q_raw)
            except ValueError as exc:
                raise ProjectValidationError(
                    f"参考答案表头包含非法题号: {q_raw}",
                    code="invalid_question_number",
                ) from exc

            answer = str(answer_raw).strip()
            q_type = classify_question(q_num, layout)
            if q_type is None:
                raise ProjectValidationError(
                    f"参考答案包含布局未声明的题号: 第 {q_num} 题",
                    code="unknown_question_number",
                )
            options = _options_for(q_type, layout)
            if options is not None and answer not in options:
                allowed = "/".join(options)
                raise ProjectValidationError(
                    f"参考答案与答题卡布局不一致。第 {q_num} 题答案为 {answer}，"
                    f"但布局声明为 {TYPE_LABELS.get(q_type, q_type)}，允许答案为 {allowed}。",
                    code="answer_layout_mismatch",
                )
    finally:
        wb.close()
