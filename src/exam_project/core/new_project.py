from __future__ import annotations

import json
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

import openpyxl

from exam_project.core.checksum import build_checksums
from exam_project.core.errors import ProjectValidationError
from exam_project.core.manifest import ProjectManifest
from exam_project.core.package import ExamProjectPackage
from exam_project.core.validation import validate_project


def _now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def _write_blank_answers(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="学号")
    ws.cell(row=2, column=1, value="参考答案")
    wb.save(path)
    wb.close()


def _blank_answer_sheet_design(name: str, student_id_digits: int) -> dict:
    return {
        "meta": {
            "title": name,
            "paper_size": "A4",
            "numbering_mode": "continuous",
        },
        "student_id": {"digit_count": student_id_digits},
        "pages": [
            {
                "sections": [
                    {
                        "type": "student_id",
                        "question_start": 0,
                        "question_count": 0,
                        "digit_count": student_id_digits,
                    },
                    {
                        "type": "choice",
                        "question_start": 1,
                        "question_count": 20,
                        "options": ["A", "B", "C", "D"],
                        "score": 3,
                    },
                ]
            },
            {
                "sections": [
                    {
                        "type": "judge",
                        "question_start": 21,
                        "question_count": 10,
                        "options": ["T", "F"],
                        "score": 2,
                    },
                    {
                        "type": "essay",
                        "question_start": 31,
                        "question_count": 1,
                        "lines_per_question": 8,
                        "score": 10,
                    },
                ]
            },
        ],
    }


def create_exam_project(
    folder_path: Path,
    *,
    name: str,
    student_id_digits: int = 10,
) -> None:
    """创建新的考试项目文件夹。

    Args:
        folder_path: 项目文件夹路径（如 Path("期中考试")）
        name: 项目名称
        student_id_digits: 学号位数
    """
    if type(student_id_digits) is not int or not (6 <= student_id_digits <= 14):
        raise ProjectValidationError(
            "学号位数必须在 6 到 14 之间",
            code="invalid_student_id_digits",
        )

    folder = folder_path.resolve()
    if folder.exists():
        raise ProjectValidationError(
            f"文件夹已存在: {folder}",
            code="folder_exists",
        )

    folder.mkdir(parents=True)
    workdir = folder
    (workdir / "design").mkdir()
    (workdir / "config").mkdir()
    (workdir / "answers").mkdir()
    (workdir / "data").mkdir()
    (workdir / "data" / "answer_sheets").mkdir()
    (workdir / "data" / "output").mkdir()
    (workdir / "data" / "processed").mkdir()

    (workdir / "design" / "answer_sheet.json").write_text(
        json.dumps(
            _blank_answer_sheet_design(name, student_id_digits),
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    (workdir / "config" / "sheet_layout.json").write_text(
        json.dumps(
            {
                "layout": {},
                "scoring": {},
                "student_id": {"digit_count": student_id_digits},
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    _write_blank_answers(workdir / "answers" / "reference_answers.xlsx")

    # 同时创建 参考答案.xlsx（供阅卷直接使用）
    _write_blank_answers(workdir / "参考答案.xlsx")

    asset_paths = [
        "design/answer_sheet.json",
        "config/sheet_layout.json",
        "answers/reference_answers.xlsx",
        "参考答案.xlsx",
    ]
    now = _now_iso()
    manifest = ProjectManifest(
        schema_version=1,
        project_id=str(uuid4()),
        name=name,
        created_at=now,
        updated_at=now,
        assets={
            "design": "design/answer_sheet.json",
            "layout": "config/sheet_layout.json",
            "answers": "answers/reference_answers.xlsx",
            "baseline": "config/blank_baseline.json",
        },
        exam={
            "student_id_digits": student_id_digits,
            "question_types": ["choice", "judge", "essay"],
        },
        checksums=build_checksums(workdir, asset_paths),
    )
    (workdir / "project.json").write_text(manifest.to_json(), encoding="utf-8")
    validate_project(workdir, manifest)

    # 生成 .examproj 归档包（放在文件夹内）
    package_path = workdir / f"{folder.name}.examproj"
    ExamProjectPackage.pack(workdir, package_path)
